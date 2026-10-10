import io
from pypdf import PdfReader
import docx

def parse_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes."""
    pdf_file = io.BytesIO(file_bytes)
    reader = PdfReader(pdf_file) 
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text.strip()

def parse_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX bytes including paragraphs and tables."""
    docx_file = io.BytesIO(file_bytes)
    doc = docx.Document(docx_file)
    full_text = []
    
    # Extract text from paragraphs
    for para in doc.paragraphs:
        if para.text:
            full_text.append(para.text)
            
    # Extract text from tables (often used in resumes)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                # To avoid duplicate content from merged cells, we clean and collect
                cell_text = " ".join([p.text.strip() for p in cell.paragraphs if p.text.strip()])
                if cell_text and cell_text not in row_text:
                    row_text.append(cell_text)
            if row_text:
                full_text.append(" | ".join(row_text))
                
    return "\n".join(full_text).strip()

def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Parses resume based on file extension."""
    ext = filename.split(".")[-1].lower()
    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext == "docx":
        return parse_docx(file_bytes)
    elif ext == "txt":
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")
